import json
from email.message import EmailMessage

import docx
import pymupdf

from copilot.parsers import parse_file
from copilot.text import normalize


def test_normalize_keeps_codes_whole():
    assert normalize("E-31 on TAC-C-0143, TS-2 is 41-2207") == "E31 on TACC0143, TS2 is 412207"
    assert normalize("re-check the take-up") == "re-check the take-up"  # plain words untouched


def test_pdf_one_chunk_per_page(tmp_path):
    path = tmp_path / "manual.pdf"
    pdf = pymupdf.open()
    for text in ["Casing bolt torque 45 N-m for all serials", "Warranty is 12 months from commissioning"]:
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(path)

    chunks = parse_file(path).chunks
    assert [c.location for c in chunks] == ["p. 1", "p. 2"]
    assert "45 N-m" in chunks[0].text


def test_scanned_pdf_page_is_ocred(tmp_path):
    # Render a text page to an image and put only the image in a new PDF: no text layer, like a scan.
    source = pymupdf.open()
    source.new_page().insert_text((72, 120), "Impeller clearance 0.38 - 0.50 mm", fontsize=20)
    image = source[0].get_pixmap(dpi=200).tobytes("png")
    scan = pymupdf.open()
    page = scan.new_page()
    page.insert_image(page.rect, stream=image)
    path = tmp_path / "scan.pdf"
    scan.save(path)

    chunks = parse_file(path).chunks
    assert chunks[0].location == "p. 1"
    assert "0.38" in chunks[0].text and "0.50" in chunks[0].text


def test_docx_sections_keep_their_tables(tmp_path):
    document = docx.Document()
    document.add_heading("6.2 Isolate", level=2)
    document.add_paragraph("Isolate the compressor.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Step 7"
    table.rows[0].cells[1].text = "Measure TP1 to TP2"
    document.add_heading("7. Records", level=1)
    document.add_paragraph("Attach Appendix A.")
    path = tmp_path / "proc.docx"
    document.save(path)

    chunks = {c.location: c.text for c in parse_file(path).chunks}
    assert "Step 7 | Measure TP1 to TP2" in chunks["6.2 Isolate"]
    assert "TP1" not in chunks["7. Records"]


def test_docx_numbered_normal_paragraph_is_a_heading(tmp_path):
    # Kyltek work instructions use Normal style for headings like "5.6 Vulling en lekcontrole [F]".
    document = docx.Document()
    document.add_paragraph("5.6 Vulling en lekcontrole [F]  (Charge and leak check [F])")
    document.add_paragraph("5.6.5\t[F] If no leak is found, top up.")
    path = tmp_path / "wi.docx"
    document.save(path)

    chunks = parse_file(path).chunks
    assert chunks[0].location.startswith("5.6 Vulling")
    assert "top up" in chunks[0].text


def test_confluence_sections_and_modified_date(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        """<html><head><title>Field : Oil Reference</title></head><body>
        <div class="page-metadata">Created by Priya, last modified on Mar 21, 2024</div>
        <div id="main-content"><div class="toc-macro"><ul><li>Intervals</li></ul></div>
        <p>Quick reference.</p>
        <h2>Intervals</h2><table><tr><th>Model</th><th>Oil change</th></tr><tr><td>VRS-75</td><td>4,000 h</td></tr></table>
        <ul><li>Never mix oils</li></ul></div></body></html>"""
    )
    parsed = parse_file(path)
    chunks = {c.location: c.text for c in parsed.chunks}
    assert parsed.note == "last modified Mar 21, 2024"
    assert "VRS-75 | 4,000 h" in chunks["Intervals"]
    assert "- Never mix oils" in chunks["Intervals"]
    assert "Intervals" not in chunks["Intro"]  # table of contents dropped


def test_email_thread_split_per_message(tmp_path):
    message = EmailMessage()
    message["From"] = "Henrik <h@belton.co.uk>"
    message["Date"] = "Mon, 11 May 2026 08:14:22 +0000"
    message["Subject"] = "RE: BX-12 tracking"
    message.set_content(
        "Replace the idler with 77-0419.\n\n________________________________\n"
        "From: Olu <o@corvane.com>\nSent: 08 May 2026 15:52\n\nBelt runs left."
    )
    path = tmp_path / "thread.eml"
    path.write_bytes(bytes(message))

    chunks = parse_file(path).chunks
    assert len(chunks) == 2
    assert chunks[0].location.startswith("Email from Henrik") and "77-0419" in chunks[0].text
    assert chunks[1].location == "Email from Olu <o@corvane.com>, 08 May 2026 15:52"


def test_teams_export_grouped_by_thread_without_deleted(tmp_path):
    def msg(id, reply_to, user, text, deleted=None):
        return {"id": id, "replyToId": reply_to, "messageType": "message", "createdDateTime": "2026-07-02T13:41:00Z",
                "deletedDateTime": deleted, "from": {"user": {"displayName": user}}, "body": {"content": f"<p>{text}</p>"}}

    path = tmp_path / "teams.json"
    path.write_text(json.dumps({"channel": {"displayName": "x"}, "value": [
        msg("1", None, "Priya", "0142 tripped E-14"),
        msg("2", "1", "Kenji", "Check firmware"),
        msg("3", "1", "Priya", "oops", deleted="2026-07-02T14:00:00Z"),
    ]}))
    chunks = parse_file(path).chunks
    assert len(chunks) == 1
    assert "Kenji: Check firmware" in chunks[0].text and "oops" not in chunks[0].text


def test_salesforce_cases_one_chunk_per_case(tmp_path):
    path = tmp_path / "cases.json"
    path.write_text(json.dumps({"records": [{
        "attributes": {"type": "Case"}, "CaseNumber": "00041988", "Asset__r": {"attributes": {}, "Name": "FRS-P-0031"},
        "Resolution__c": "Torqued to 38 N-m", "Parts_Used__c": None,
        "CaseComments": {"totalSize": 1, "records": [{"CommentBody": "gasket seat cracked"}]},
    }]}))
    chunk = parse_file(path).chunks[0]
    assert chunk.location == "Case 00041988"
    assert "Asset: FRS-P-0031" in chunk.text
    assert "CaseComments.CommentBody: gasket seat cracked" in chunk.text
    assert "Parts_Used" not in chunk.text and "attributes" not in chunk.text


def test_csv_semicolon_export_one_chunk_per_row(tmp_path):
    path = tmp_path / "servio.csv"
    path.write_text("werkorder;object;omschrijving\n2600411;RTM-H-0007;bijgevuld 2,5 kg\n")
    chunk = parse_file(path).chunks[0]
    assert chunk.location == "werkorder 2600411"
    assert chunk.text == "werkorder: 2600411; object: RTM-H-0007; omschrijving: bijgevuld 2,5 kg"
